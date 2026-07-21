#!/usr/bin/env python3
"""Test tech stack detection on all domains from the Excel spreadsheet."""
import asyncio
import csv
import time
import httpx
import openpyxl
from src.scraping.parser.tech_parser import TechParser

XLSX_PATH = "/Users/deep/Downloads/100-200 Test.xlsx"
OUTPUT_CSV = "/Users/deep/Downloads/tech_stack_results.csv"
CONCURRENT = 8
TIMEOUT_SEC = 20

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                  "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.5",
    "Accept-Encoding": "gzip, deflate",
    "Connection": "keep-alive",
}

semaphore = asyncio.Semaphore(CONCURRENT)


async def fetch(client: httpx.AsyncClient, domain: str) -> tuple[str, dict, str | None]:
    async with semaphore:
        url = domain if domain.startswith(("http://", "https://")) else f"https://{domain}"
        try:
            resp = await client.get(url, headers=HEADERS, follow_redirects=True, timeout=TIMEOUT_SEC)
            return resp.text, dict(resp.headers), None
        except Exception as e:
            return "", {}, repr(e)


def parse_tech(html: str, resp_headers: dict) -> dict:
    clean_headers = {str(k): str(v) for k, v in resp_headers.items()}
    parser = TechParser(html, clean_headers)
    return parser.detect()


def format_detections(result: dict) -> str:
    """Format detections as a compact string for CSV output."""
    parts = []
    for d in sorted(result.get("detections", []), key=lambda x: (-{"high": 3, "medium": 2, "low": 1}[x["confidence"]], x["name"])):
        parts.append(f"{d['name']} ({d['confidence']})")
    return "; ".join(parts)


def main():
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        domain = ws.cell(r, 2).value
        country = ws.cell(r, 3).value
        if domain:
            rows.append((name, str(domain).strip(), str(country).strip() if country else ""))

    print(f"Total domains: {len(rows)}")
    print(f"Concurrency: {CONCURRENT} | Timeout: {TIMEOUT_SEC}s")
    print(f"Output CSV: {OUTPUT_CSV}\n")

    results = []

    async def run_all():
        async with httpx.AsyncClient() as client:
            tasks = [fetch(client, domain) for _, domain, _ in rows]
            return await asyncio.gather(*tasks)

    t0 = time.time()
    fetch_results = asyncio.run(run_all())
    total_time = time.time() - t0

    # Process results
    for i, ((company, domain, country), (html, resp_headers, error)) in enumerate(zip(rows, fetch_results)):
        if error:
            results.append({
                "company": company,
                "domain": domain,
                "country": country,
                "status": "FETCH_ERROR",
                "html_size": 0,
                "error": error[:120],
                "platform": "",
                "detection_count": 0,
                "detections": "",
                "high_confidence": "",
                "medium_confidence": "",
                "low_confidence": "",
            })
            continue

        if not html or len(html) < 200:
            results.append({
                "company": company,
                "domain": domain,
                "country": country,
                "status": "EMPTY",
                "html_size": len(html),
                "error": "",
                "platform": "",
                "detection_count": 0,
                "detections": "",
                "high_confidence": "",
                "medium_confidence": "",
                "low_confidence": "",
            })
            continue

        tech_result = parse_tech(html, resp_headers)
        detections = tech_result.get("detections", [])
        high = [d["name"] for d in detections if d["confidence"] == "high"]
        med = [d["name"] for d in detections if d["confidence"] == "medium"]
        low = [d["name"] for d in detections if d["confidence"] == "low"]

        results.append({
            "company": company,
            "domain": domain,
            "country": country,
            "status": "OK",
            "html_size": len(html),
            "error": "",
            "platform": tech_result.get("platform") or "",
            "detection_count": len(detections),
            "detections": format_detections(tech_result),
            "high_confidence": ", ".join(high),
            "medium_confidence": ", ".join(med),
            "low_confidence": ", ".join(low),
        })

    # Write CSV
    fieldnames = ["company", "domain", "country", "status", "html_size", "error",
                  "platform", "detection_count", "detections",
                  "high_confidence", "medium_confidence", "low_confidence"]
    with open(OUTPUT_CSV, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)

    # Print summary
    ok_count = sum(1 for r in results if r["status"] == "OK")
    fetch_errors = sum(1 for r in results if r["status"] == "FETCH_ERROR")
    empty = sum(1 for r in results if r["status"] == "EMPTY")
    with_platform = sum(1 for r in results if r["platform"])

    print("=" * 60)
    print("SUMMARY")
    print("=" * 60)
    print(f"Total domains:    {len(results)}")
    print(f"Successful:       {ok_count}")
    print(f"Fetch errors:     {fetch_errors}")
    print(f"Empty responses:  {empty}")
    print(f"Platform detected:{with_platform}/{ok_count}")
    print(f"Total time:       {total_time:.1f}s (avg {total_time/len(results):.1f}s/domain)")
    print()

    # Confidence distribution
    all_dets = []
    for r in results:
        if r["status"] == "OK":
            all_dets.extend([d.strip() for d in r["detections"].split(";") if d.strip()])

    high_count = sum(1 for d in all_dets if "(high)" in d)
    med_count = sum(1 for d in all_dets if "(medium)" in d)
    low_count = sum(1 for d in all_dets if "(low)" in d)

    print(f"Total detections: {len(all_dets)}")
    print(f"  HIGH:   {high_count}")
    print(f"  MEDIUM: {med_count}")
    print(f"  LOW:    {low_count}")
    print()

    # Top detected technologies
    from collections import Counter
    name_counts = Counter()
    for d in all_dets:
        name = d.rsplit(" (", 1)[0] if " (" in d else d
        name_counts[name] += 1
    print("Top 15 detected technologies:")
    for name, count in name_counts.most_common(15):
        print(f"  {name:<30} {count:>3} sites")
    print()

    # Platform breakdown
    plat_counts = Counter(r["platform"] for r in results if r["platform"])
    print("Platforms detected:")
    for name, count in plat_counts.most_common():
        print(f"  {name:<30} {count:>3} sites")
    print()

    # Fetch errors detail
    if fetch_errors:
        print("Fetch errors:")
        for r in results:
            if r["status"] == "FETCH_ERROR":
                print(f"  {r['domain']:<40} {r['error']}")
        print()

    print(f"Results saved to: {OUTPUT_CSV}")


if __name__ == "__main__":
    main()
